"""Results export (task 4.5): an Excel workbook that keeps Bengali names intact, and the data behind the
printable result cards."""
import uuid
from datetime import datetime, timedelta, timezone
from io import BytesIO

import openpyxl

from app import models
from tests.test_organizations import headers_for, make_org, make_user


def setup(db, show_rank=False):
    org = make_org(db, f"Export {uuid.uuid4().hex[:6]}")
    org.timezone = "Asia/Dhaka"
    owner = make_user(db, f"owner-{uuid.uuid4().hex[:6]}@export.com", org, models.ROLE_OWNER)
    now = datetime.now(timezone.utc)
    q1 = models.Question(organization_id=org.id, title="2 + 2 = ?", complexity="easy", type="single_choice",
                         options=["3", "4"], correct_answers=["4"], max_score=1, tags=[])
    q2 = models.Question(organization_id=org.id, title="Pick the primes", complexity="easy", type="multi_choice",
                         options=["2", "4", "5"], correct_answers=["2", "5"], max_score=2, tags=[])
    q3 = models.Question(organization_id=org.id, title="Explain inertia", complexity="medium", type="text",
                         options=None, correct_answers=None, max_score=5, tags=[])
    exam = models.Exam(organization_id=org.id, title="ভৌতবিজ্ঞান মডেল টেস্ট", start_time=now - timedelta(hours=3),
                       end_time=now - timedelta(minutes=1), duration_minutes=60, is_published=True, show_rank=show_rank)
    exam.questions = [q1, q2, q3]
    db.add_all([q1, q2, q3, exam])
    db.commit()
    return org, owner, exam, (q1, q2, q3)


def attempt(db, org, exam, name, answers, minutes, graded=None, finished=True):
    student = make_user(db, f"s-{uuid.uuid4().hex[:6]}@export.com", org, models.ROLE_CANDIDATE)
    student.full_name = name
    start = datetime.now(timezone.utc) - timedelta(hours=2)
    row = models.ExamAttempt(exam_id=exam.id, student_id=student.id, start_time=start)
    db.add(row)
    db.commit()
    for question, value in answers.items():
        db.add(models.Answer(attempt_id=row.id, question_id=question.id, answer_data=value))
    db.commit()
    if finished:
        from app import crud
        crud.finish_attempt(db, row, ended_at=start + timedelta(minutes=minutes))
    if graded is not None:
        answer = db.query(models.Answer).filter_by(attempt_id=row.id, question_id=graded[0].id).one()
        db.add(models.Evaluation(answer_id=answer.id, evaluated_by=graded[2].id, score_awarded=graded[1]))
        db.commit()
    return student, row


def workbook(client, owner, org, exam):
    res = client.get(f"/admin/exams/{exam.id}/export.xlsx", headers=headers_for(owner, org))
    assert res.status_code == 200, res.text
    assert res.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert "filename*=UTF-8''" in res.headers["content-disposition"]
    return openpyxl.load_workbook(BytesIO(res.content))


def test_the_workbook_keeps_bengali_names_and_ranks_the_students(client, test_db):
    org, owner, exam, (q1, q2, q3) = setup(test_db)
    attempt(test_db, org, exam, "মিম আক্তার", {q1: ["4"], q2: ["5", "2"], q3: "Mass resists change"}, 30, graded=(q3, 4, owner))
    attempt(test_db, org, exam, "Rafi Hasan", {q1: ["3"], q2: ["2", "5"]}, 20)
    attempt(test_db, org, exam, "Still writing", {q1: ["4"]}, 0, finished=False)

    wb = workbook(client, owner, org, exam)
    assert wb.sheetnames == ["Results", "Marks by question", "Questions", "Exam"]
    results = wb["Results"]
    rows = [[c.value for c in r] for r in results.iter_rows(min_row=2)]
    assert [(r[0], r[1], r[3], r[4]) for r in rows] == [(1, "মিম আক্তার", 7, 8), (2, "Rafi Hasan", 2, 8), (None, "Still writing", None, None)]
    assert rows[0][9] == "Submitted" and rows[2][9] == "In progress"
    assert rows[0][6] == timedelta(minutes=30)  # time taken, stored as an Excel duration

    marks = [[c.value for c in r] for r in wb["Marks by question"].iter_rows(values_only=False)]
    assert marks[0][2:] == ["Q1 (1)", "Q2 (2)", "Q3 (5)"]
    assert marks[1][2:] == [1, 2, 4]
    assert marks[2][2:] == [0, 2, 0]  # wrong single answer, right multiple answer, written question left blank
    assert wb["Exam"]["B1"].value == "ভৌতবিজ্ঞান মডেল টেস্ট"
    assert wb["Exam"]["B10"].value == "Asia/Dhaka"


def test_an_ungraded_written_answer_says_so(client, test_db):
    org, owner, exam, (q1, q2, q3) = setup(test_db)
    attempt(test_db, org, exam, "Waiting", {q3: "An answer nobody marked yet"}, 10)
    marks = list(workbook(client, owner, org, exam)["Marks by question"].iter_rows(min_row=2, values_only=True))
    assert marks[0][2:] == (0, 0, "Not graded")


def test_a_name_that_looks_like_a_formula_stays_text(client, test_db):
    org, owner, exam, (q1, _, _) = setup(test_db)
    attempt(test_db, org, exam, '=HYPERLINK("http://evil.example","click")', {q1: ["4"]}, 10)
    cell = workbook(client, owner, org, exam)["Results"]["B2"]
    assert cell.value == '=HYPERLINK("http://evil.example","click")'
    assert cell.data_type == "s"


def test_another_organization_gets_nothing_and_students_are_refused(client, test_db):
    org, owner, exam, (q1, _, _) = setup(test_db)
    student, _ = attempt(test_db, org, exam, "Mim", {q1: ["4"]}, 10)
    other = make_org(test_db, "Elsewhere")
    stranger = make_user(test_db, f"x-{uuid.uuid4().hex[:6]}@else.com", other, models.ROLE_OWNER)
    for path in ("export.xlsx", "result-cards"):
        assert client.get(f"/admin/exams/{exam.id}/{path}", headers=headers_for(stranger, other)).status_code == 404
        assert client.get(f"/admin/exams/{exam.id}/{path}", headers=headers_for(student, org)).status_code == 403


def test_result_cards_carry_marks_by_question_and_a_rank_only_when_the_exam_shows_it(client, test_db):
    org, owner, exam, (q1, q2, q3) = setup(test_db)
    attempt(test_db, org, exam, "First", {q1: ["4"], q2: ["2", "5"]}, 15)
    _, second = attempt(test_db, org, exam, "Second", {q1: ["4"]}, 25)
    attempt(test_db, org, exam, "Unfinished", {q1: ["4"]}, 0, finished=False)

    data = client.get(f"/admin/exams/{exam.id}/result-cards", headers=headers_for(owner, org)).json()
    assert data["submitted"] == 2 and len(data["cards"]) == 2
    assert data["exam"]["start_time"].endswith("+00:00") and data["cards"][0]["submitted_at"].endswith("+00:00")
    first = data["cards"][0]
    assert first["student"]["full_name"] == "First" and first["rank"] is None
    assert (first["score"], first["total"], first["percentage"]) == (3, 8, 37.5)
    assert [(q["n"], q["worth"], q["earned"]) for q in first["questions"]] == [(1, 1, 1), (2, 2, 2), (3, 5, 0)]
    assert "violations" not in first  # cards go to students: no integrity details

    exam.show_rank = True
    test_db.commit()
    one = client.get(f"/admin/exams/{exam.id}/result-cards", params={"attempt": str(second.id)}, headers=headers_for(owner, org)).json()
    assert [c["student"]["full_name"] for c in one["cards"]] == ["Second"]
    assert one["cards"][0]["rank"] == 2 and one["submitted"] == 2
