"""Random question pools (task 4.6, PRD US17/AC18): each attempt draws its own questions once, stores them,
and keeps them through reloads and changes to the bank; grading and everything after use the stored paper."""
import uuid
from datetime import datetime, timedelta, timezone

from app import crud, models
from tests.test_organizations import headers_for, make_org, make_user


def question(db, org, title, complexity="easy", tags=("physics",), marks=3):
    q = models.Question(organization_id=org.id, title=title, complexity=complexity, type="single_choice",
                        options=["right", "wrong"], correct_answers=["right"], max_score=marks, tags=list(tags))
    db.add(q)
    db.commit()
    return q


def setup(db):
    org = make_org(db, f"Pools {uuid.uuid4().hex[:6]}")
    owner = make_user(db, f"owner-{uuid.uuid4().hex[:6]}@pools.com", org, models.ROLE_OWNER)
    fixed = question(db, org, "Fixed: units of force", tags=("mechanics",), marks=2)
    physics_easy = [question(db, org, f"Easy physics {i}") for i in range(8)]
    hard = [question(db, org, f"Hard physics {i}", complexity="hard") for i in range(3)]
    chemistry = [question(db, org, f"Easy chemistry {i}", tags=("chemistry",)) for i in range(4)]
    return org, owner, fixed, physics_easy, hard, chemistry


def exam_body(fixed_ids, rules, **extra):
    now = datetime.now(timezone.utc)
    return {"title": "Physics pool test", "start_time": (now - timedelta(minutes=5)).isoformat(),
            "end_time": (now + timedelta(hours=2)).isoformat(), "duration_minutes": 30,
            "question_ids": [str(i) for i in fixed_ids], "pool_rules": rules, "require_fullscreen": False, **extra}


def create(client, owner, org, body):
    res = client.post("/admin/exams/", json=body, headers=headers_for(owner, org))
    assert res.status_code == 200, res.text
    exam = res.json()
    assert client.post(f"/admin/exams/{exam['id']}/publish", headers=headers_for(owner, org)).status_code == 200
    return exam


def student(db, org):
    return make_user(db, f"s-{uuid.uuid4().hex[:6]}@pools.com", org, models.ROLE_CANDIDATE)


RULE = {"count": 5, "complexity": "easy", "tag": "Physics", "marks": 1}


def test_ac18_five_random_easy_physics_questions_stay_the_same_after_a_reload_and_bank_changes(client, test_db):
    org, owner, fixed, physics_easy, hard, chemistry = setup(test_db)
    exam = create(client, owner, org, exam_body([fixed.id], [RULE]))
    assert exam["pool_rules"][0]["tag"] == "physics"  # stored the way the bank writes tags
    assert exam["question_count"] == 6 and exam["total_marks"] == 2 + 5 * 1

    kid = student(test_db, org)
    h = headers_for(kid, org)
    session = client.post(f"/student/exams/{exam['id']}/start", headers=h).json()
    ids = {q["id"] for q in session["exam"]["questions"]}
    drawn = ids - {str(fixed.id)}
    assert str(fixed.id) in ids and len(ids) == 6 and len(drawn) == 5
    assert drawn <= {str(q.id) for q in physics_easy}  # only easy + physics
    assert all(q["max_score"] == 1 for q in session["exam"]["questions"] if q["id"] in drawn)

    # The bank changes: two drawn questions archived, three new matching ones added
    for q in test_db.query(models.Question).filter(models.Question.id.in_([uuid.UUID(i) for i in list(drawn)[:2]])):
        q.archived_at = datetime.now(timezone.utc)
    for i in range(3):
        question(test_db, org, f"New easy physics {i}")
    test_db.commit()

    attempt_id = session["attempt"]["id"]
    again = client.post(f"/student/attempts/{attempt_id}/resume", headers=h).json()
    assert {q["id"] for q in again["exam"]["questions"]} == ids
    restart = client.post(f"/student/exams/{exam['id']}/start", headers=h).json()
    assert {q["id"] for q in restart["exam"]["questions"]} == ids

    # Grading uses the stored paper and its frozen marks
    for q in again["exam"]["questions"]:
        res = client.post(f"/student/attempts/{attempt_id}/save-answer", headers=h, json={"question_id": q["id"], "answer_data": ["right"]})
        assert res.status_code == 200, res.text
    client.post(f"/student/attempts/{attempt_id}/submit", headers=h)
    attempt = test_db.get(models.ExamAttempt, uuid.UUID(attempt_id))
    test_db.refresh(attempt)
    assert (attempt.score, attempt.total_possible_score) == (7, 7)


def test_every_student_draws_their_own_set(client, test_db):
    org, owner, fixed, physics_easy, _, _ = setup(test_db)
    exam = create(client, owner, org, exam_body([], [RULE]))
    papers = set()
    for _ in range(4):
        session = client.post(f"/student/exams/{exam['id']}/start", headers=headers_for(student(test_db, org), org)).json()
        papers.add(frozenset(q["id"] for q in session["exam"]["questions"]))
    assert len(papers) > 1  # 56 possible sets of 5 from 8: four students all alike is vanishingly unlikely


def test_a_question_not_on_the_students_paper_is_refused(client, test_db):
    org, owner, fixed, physics_easy, _, _ = setup(test_db)
    exam = create(client, owner, org, exam_body([], [RULE]))
    kid = student(test_db, org)
    h = headers_for(kid, org)
    session = client.post(f"/student/exams/{exam['id']}/start", headers=h).json()
    mine = {q["id"] for q in session["exam"]["questions"]}
    other = next(q for q in physics_easy if str(q.id) not in mine)
    res = client.post(f"/student/attempts/{session['attempt']['id']}/save-answer", headers=h,
                      json={"question_id": str(other.id), "answer_data": ["right"]})
    assert res.status_code == 403 and res.json()["detail"] == "Question does not belong to this exam"


def test_rules_the_bank_cannot_fill_are_refused(client, test_db):
    org, owner, fixed, *_ = setup(test_db)
    h = headers_for(owner, org)
    too_many = client.post("/admin/exams/", headers=h, json=exam_body([], [{**RULE, "count": 9}]))
    assert too_many.status_code == 422 and "Only 8 questions" in too_many.json()["detail"]
    hard_only = client.post("/admin/exams/", headers=h, json=exam_body([], [{"count": 4, "complexity": "hard", "tag": "physics"}]))
    assert hard_only.status_code == 422
    overlap = client.post("/admin/exams/", headers=h, json=exam_body([], [{**RULE, "count": 6}, {"count": 6, "tag": "physics", "complexity": "easy"}]))
    assert overlap.status_code == 422 and "overlap" in overlap.json()["detail"]
    nothing = client.post("/admin/exams/", headers=h, json=exam_body([], []))
    assert nothing.status_code == 422


def test_a_drawn_question_is_archived_not_deleted(client, test_db):
    org, owner, fixed, physics_easy, _, _ = setup(test_db)
    exam = create(client, owner, org, exam_body([], [RULE]))
    session = client.post(f"/student/exams/{exam['id']}/start", headers=headers_for(student(test_db, org), org)).json()
    drawn = uuid.UUID(session["exam"]["questions"][0]["id"])
    assert crud.question_in_use(test_db, drawn)
    undrawn = next(q for q in physics_easy if str(q.id) not in {x["id"] for x in session["exam"]["questions"]})
    assert not crud.question_in_use(test_db, undrawn.id)


def test_students_see_the_full_count_and_publishing_checks_the_bank_again(client, test_db):
    org, owner, fixed, physics_easy, _, _ = setup(test_db)
    exam = create(client, owner, org, exam_body([fixed.id], [RULE]))
    listing = client.get("/student/exams/", headers=headers_for(student(test_db, org), org)).json()
    assert next(e for e in listing if e["id"] == exam["id"])["question_count"] == 6

    h = headers_for(owner, org)
    assert client.post(f"/admin/exams/{exam['id']}/unpublish", headers=h).status_code == 200
    for q in physics_easy[:4]:
        q.archived_at = datetime.now(timezone.utc)
    test_db.commit()
    res = client.post(f"/admin/exams/{exam['id']}/publish", headers=h)
    assert res.status_code == 409 and "Only 4 questions" in res.json()["detail"]


def test_the_export_labels_drawn_slots_and_lists_the_rules(client, test_db):
    import openpyxl
    from io import BytesIO

    org, owner, fixed, physics_easy, _, _ = setup(test_db)
    exam = create(client, owner, org, exam_body([fixed.id], [RULE]))
    kid = student(test_db, org)
    session = client.post(f"/student/exams/{exam['id']}/start", headers=headers_for(kid, org)).json()
    client.post(f"/student/attempts/{session['attempt']['id']}/submit", headers=headers_for(kid, org))

    wb = openpyxl.load_workbook(BytesIO(client.get(f"/admin/exams/{exam['id']}/export.xlsx", headers=headers_for(owner, org)).content))
    header = [c.value for c in next(wb["Marks by question"].iter_rows(max_row=1))]
    assert header == ["Name", "Email", "Q1 (2)"] + [f"Q{i} (1, random)" for i in range(2, 7)]
    rules = list(wb["Questions"].iter_rows(min_row=3, values_only=True))
    assert rules == [("Q2-Q6", '5 random easy questions tagged "physics"', "Drawn per student", 1)]
    assert wb["Exam"]["B6"].value == 7  # total marks: 2 fixed + 5 x 1
